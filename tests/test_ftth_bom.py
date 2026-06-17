import unittest
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
INPUTS = ROOT / "inputs"
BELL = next(INPUTS.glob("*BELL*.xlsx"), None)
CATALOG = next(INPUTS.glob("KATALOGI*.xlsx"), None)
JEDLNIA = next(INPUTS.glob("*Jedlnia*.gpkg"), None)
DEBE = INPUTS / "PW_Debe_Wielkie_OPP25.gpkg"
DPW = INPUTS / "DPW_F03013604.gpkg"
DPW_13171 = INPUTS / "DPW_F03013171.gpkg"
ZAD2 = ROOT / "przyklady" / "zad 2" / "Nadanie_F03008070_v2.gpkg"


@unittest.skipUnless(
    all(path is not None and path.exists() for path in [BELL, CATALOG, JEDLNIA, DEBE, DPW, DPW_13171, ZAD2]),
    "Pominieto testy integracyjne: brak lokalnych plikow GPKG/XLSX w inputs/przyklady.",
)
class FtthBomTests(unittest.TestCase):
    def test_bell_history_loads_rows_and_classifies_jedlnia_materials(self):
        from ftth_bom.bell_history import load_bell_history

        rows = load_bell_history(BELL)
        self.assertGreater(len(rows), 9000)

        jedlnia_rows = [row for row in rows if "F/03013155" in row.all_fx]
        self.assertGreater(len(jedlnia_rows), 20)
        self.assertTrue(
            any(row.sap == "2200003264" and row.material_class == "DAC_2J" for row in jedlnia_rows)
        )

    def test_gpkg_reader_finds_project_keys_and_unique_cable_metrics(self):
        from ftth_bom.gpkg_reader import read_gpkg_project

        jedlnia = read_gpkg_project(JEDLNIA)
        debe = read_gpkg_project(DEBE)

        self.assertIn("F/03013155", jedlnia.f_numbers)
        self.assertIn("F/03008068", debe.f_numbers)
        self.assertNotIn("F/03008066", debe.f_numbers)
        self.assertIn("F/03008066", debe.metadata["related_f_numbers"])
        self.assertGreater(jedlnia.hh_total, 0)
        self.assertGreater(debe.hh_total, 0)
        self.assertGreater(len(jedlnia.cable_summary), 0)
        self.assertEqual(
            len({c.cable_id for c in jedlnia.cable_edges}),
            len(jedlnia.cable_edges),
        )

    def test_gpkg_reader_scans_project_layers_for_material_takeoff(self):
        from ftth_bom.gpkg_reader import read_gpkg_project

        project = read_gpkg_project(JEDLNIA)
        takeoff = {item.material_class: item for item in project.material_takeoff}

        self.assertEqual(takeoff["MIKRORURKA_12_8"].length_m, 1155)
        self.assertEqual(takeoff["MIKRORURKA_14_10"].length_m, 308)
        self.assertEqual(takeoff["PSB_H_144"].count, 1)
        self.assertEqual(takeoff["SUS_PH_S"].count, 4)
        self.assertEqual(takeoff["MUFA_SSC2110"].count, 1)
        self.assertEqual(takeoff["SPLITTER_1X64"].count, 1)

    def test_material_classifier_handles_adss_splitters_and_mikrokable(self):
        from ftth_bom.rules import classify_material_name

        self.assertEqual(classify_material_name("ADSS 12J"), "ADSS_12J")
        self.assertEqual(classify_material_name("ADSS 24J"), "ADSS_24J")
        self.assertEqual(classify_material_name("S-PL-108-TUBE-900-SCA"), "SPLITTER")
        self.assertEqual(classify_material_name("MI-MKF 12J"), "MIKROKABEL_12J")
        self.assertEqual(classify_material_name("MI-MKF 36J"), "MIKROKABEL_36J")

    def test_material_classifier_maps_dac_4j_to_dac_6j_substitute(self):
        from ftth_bom.rules import classify_material_name

        self.assertEqual(classify_material_name("DAC 4J"), "DAC_6J")

    def test_material_classifier_keeps_dac_12j_separate_from_dac_2j(self):
        from ftth_bom.rules import classify_material_name

        self.assertEqual(classify_material_name("DAC 12J"), "DAC_12J")

    def test_material_classifier_handles_project_layer_materials(self):
        from ftth_bom.rules import classify_material_name

        self.assertEqual(classify_material_name("FP-MR-G-12/8 (pomaranczowy)"), "MIKRORURKA_12_8")
        self.assertEqual(classify_material_name("FP-MR-G-14/10 (pomaranczowy)"), "MIKRORURKA_14_10")
        self.assertEqual(classify_material_name("RURA HDPE-UV FI 40/3,7"), "HDPE_UV_40")
        self.assertEqual(classify_material_name("UCHWYT DYSTANS.RUR 25-50 NA SL.ZELBET"), "DYSTANS_HDPE_UV")
        self.assertEqual(classify_material_name("TASMA STALOWA TSM/10-07-J"), "TASMA_STALOWA_10MM")
        self.assertEqual(classify_material_name("TASMA STALOWA TSM/20-07-J"), "TASMA_STALOWA")
        self.assertEqual(classify_material_name("PSB-H-144-GM"), "PSB_H_144")
        self.assertEqual(classify_material_name("SUS-PH-S"), "SUS_PH_S")
        self.assertEqual(classify_material_name("SSC2110-FM-48"), "MUFA_SSC2110")
        self.assertEqual(classify_material_name("SPL1x64/1216/SCA"), "SPLITTER_1X64")

    def test_aerial_poles_are_matched_from_geometry_without_concept_layers(self):
        from ftth_bom.gpkg_reader import read_gpkg_project

        project = read_gpkg_project(ZAD2)

        self.assertEqual(len(project.used_poles), 21)
        self.assertEqual(sum(1 for pole in project.used_poles if pole.source_layer == "Plan_Obiekty"), 4)
        self.assertFalse(any(pole.source_layer.startswith("K ") for pole in project.used_poles))
        self.assertTrue(all(pole.matched_cables for pole in project.used_poles))

    def test_reader_does_not_treat_sus_posts_as_aerial_poles(self):
        from ftth_bom.gpkg_reader import read_gpkg_project

        project = read_gpkg_project(JEDLNIA)

        self.assertEqual(len(project.used_poles), 18)
        self.assertFalse(any("SUS" in pole.model.upper() for pole in project.used_poles))
        self.assertFalse(any("SUS" in pole.pole_id.upper() for pole in project.used_poles))

    def test_reader_counts_used_energy_poles_for_nameplate_tape(self):
        from ftth_bom.gpkg_reader import read_gpkg_project

        debe = read_gpkg_project(DEBE)
        jedlnia = read_gpkg_project(JEDLNIA)
        dpw = read_gpkg_project(DPW_13171)

        self.assertEqual(debe.metadata["used_energy_poles"], 23)
        self.assertEqual(jedlnia.metadata["used_energy_poles"], 18)
        self.assertEqual(dpw.metadata["used_energy_poles"], 7)

    def test_preferences_default_to_learned_material_class_not_current_fx(self):
        from ftth_bom.bell_history import load_bell_history
        from ftth_bom.preferences import PreferenceStore

        store = PreferenceStore.from_history(load_bell_history(BELL))
        recommendation = store.recommend("DAC_2J", f_numbers=["F/03013155"], project_tokens=["OPP 0025"])

        self.assertIsNotNone(recommendation)
        self.assertEqual(recommendation.sap, "2200003264")
        self.assertEqual(recommendation.match_level, "learned_class")
        self.assertGreater(recommendation.score, 0)

    def test_oap_boxes_are_learned_separately_from_oap_reserve_frames(self):
        from ftth_bom.bell_history import load_bell_history
        from ftth_bom.preferences import PreferenceStore

        store = PreferenceStore.from_history(load_bell_history(BELL))
        recommendation = store.recommend("OAP_48")

        self.assertIsNotNone(recommendation)
        self.assertEqual(recommendation.sap, "2200023907")
        self.assertIn("OAP-48", recommendation.name)

    def test_engine_generates_reviewable_bom_and_sap_view(self):
        from ftth_bom.engine import run_analysis

        result = run_analysis(
            gpkg_path=JEDLNIA,
            catalog_path=CATALOG,
            bell_path=BELL,
            task_name="test_jedlnia",
        )

        self.assertGreater(len(result.order_rows), 5)
        self.assertGreater(len(result.sap_rows), 0)
        self.assertTrue(any(row.sap == "2200003264" for row in result.order_rows))
        self.assertTrue(any(row.sap == "2200003277" and row.unit.lower() == "m" for row in result.order_rows))
        self.assertFalse(any(row.match_level == "exact_f" for row in result.order_rows))
        self.assertTrue(any(item.topic == "Historia BELL" for item in result.issues))

    def test_engine_skips_adss_2j_service_drops_from_jedlnia_order(self):
        from ftth_bom.engine import run_analysis

        result = run_analysis(
            gpkg_path=JEDLNIA,
            catalog_path=CATALOG,
            bell_path=BELL,
            task_name="test_jedlnia_service_drops",
        )

        self.assertFalse(any(row.material_class == "ADSS_2J" for row in result.order_rows))

    def test_engine_hides_non_orderable_service_drops_from_result_preferences(self):
        from ftth_bom.engine import run_analysis

        result = run_analysis(
            gpkg_path=JEDLNIA,
            catalog_path=CATALOG,
            bell_path=BELL,
            task_name="test_jedlnia_preferences_without_service_drops",
        )

        self.assertFalse(any(pref.material_class == "ADSS_2J" for pref in result.preferences))

    def test_engine_orders_materials_from_project_layer_takeoff(self):
        from ftth_bom.engine import run_analysis

        result = run_analysis(
            gpkg_path=JEDLNIA,
            catalog_path=CATALOG,
            bell_path=BELL,
            task_name="test_jedlnia_project_layer_takeoff",
        )
        rows = {row.material_class: row for row in result.order_rows}

        self.assertEqual(rows["MIKRORURKA_12_8"].sap, "2200004586")
        self.assertEqual(rows["MIKRORURKA_12_8"].qty_pw, "1155 m")
        self.assertEqual(rows["MIKRORURKA_14_10"].sap, "2200004524")
        self.assertEqual(rows["MIKRORURKA_14_10"].qty_pw, "308 m")
        self.assertEqual(rows["PSB_H_144"].sap, "2200023135")
        self.assertEqual(rows["SUS_PH_S"].sap, "2200023700")
        self.assertEqual(rows["SUS_PH_S"].qty_order, 4)
        self.assertEqual(rows["MUFA_SSC2110"].sap, "2200028376")
        self.assertEqual(rows["SPLITTER_1X64"].sap, "2200006269")

    def test_engine_uses_geometry_matched_poles_for_aerial_dead_end_hardware(self):
        from ftth_bom.engine import run_analysis

        result = run_analysis(
            gpkg_path=DEBE,
            catalog_path=CATALOG,
            bell_path=BELL,
            task_name="test_debe_poles",
        )

        rows = [row for row in result.order_rows if row.material_class == "UCHWYT_ODCIAGOWY"]
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].qty_order, 64)
        self.assertIn("geometr", rows[0].basis.lower())

    def test_engine_applies_revised_aerial_and_pigtail_rules_for_debe(self):
        from ftth_bom.engine import run_analysis

        result = run_analysis(
            gpkg_path=DEBE,
            catalog_path=CATALOG,
            bell_path=BELL,
            task_name="test_debe_accessories",
        )
        rows = {row.material_class: row for row in result.order_rows}

        self.assertEqual(result.project.metadata["pigtail_splices"], 20)
        self.assertEqual(result.project.metadata["splice_sleeves"], 38)
        self.assertEqual(result.project.metadata["oap_on_used_poles"], 10)
        self.assertEqual(result.project.metadata["oap_on_energy_poles"], 8)
        self.assertEqual(rows["PIGTAIL_SC_APC"].qty_order, 20)
        self.assertEqual(rows["OSLONKA_SPAWU"].qty_order, 38)
        self.assertNotIn("ADAPTER_SC_APC", rows)
        self.assertNotIn("ZAMEK_ABLOY", rows)
        self.assertEqual(rows["UCHWYT_ODCIAGOWY"].qty_order, 64)
        self.assertEqual(rows["HAK_UNIWERSALNY"].sap, "2200003716")
        self.assertEqual(rows["HAK_UNIWERSALNY"].qty_order, 64)
        self.assertEqual(rows["TASMA_STALOWA"].sap, "2200003758")
        self.assertEqual(rows["TASMA_STALOWA"].unit, "ROL")
        self.assertEqual(rows["TASMA_STALOWA"].qty_order, 3)
        self.assertIn("126", rows["TASMA_STALOWA"].qty_pw)
        self.assertEqual(rows["TASMA_STALOWA_10MM"].sap, "2200003733")
        self.assertEqual(rows["TASMA_STALOWA_10MM"].qty_order, 1)
        self.assertIn("11.2 m", rows["TASMA_STALOWA_10MM"].qty_pw)
        self.assertIn("23 slupow EN", rows["TASMA_STALOWA_10MM"].qty_pw)
        self.assertEqual(rows["KLAMRA_TASMY"].sap, "2200003756")
        self.assertEqual(rows["KLAMRA_TASMY"].unit, "PAK")
        self.assertEqual(rows["KLAMRA_TASMY"].qty_order, 2)
        self.assertIn("173", rows["KLAMRA_TASMY"].qty_pw)
        self.assertEqual(rows["DYSTANS_OAP"].sap, "2200028041")
        self.assertEqual(rows["DYSTANS_OAP"].qty_order, 8)

    def test_engine_counts_unique_pigtail_ports_and_real_splice_sleeves_for_jedlnia(self):
        from ftth_bom.engine import run_analysis

        result = run_analysis(
            gpkg_path=JEDLNIA,
            catalog_path=CATALOG,
            bell_path=BELL,
            task_name="test_jedlnia_real_pigtails",
        )
        rows = {row.material_class: row for row in result.order_rows}

        self.assertEqual(result.project.metadata["pigtail_splices_raw"], 117)
        self.assertEqual(result.project.metadata["pigtail_splices"], 115)
        self.assertEqual(result.project.metadata["fiber_splice_connections"], 97)
        self.assertEqual(result.project.metadata["splice_sleeves"], 214)
        self.assertEqual(result.project.metadata["adapter_required_pigtails"], 98)
        self.assertEqual(rows["PIGTAIL_SC_APC"].qty_order, 115)
        self.assertEqual(rows["ADAPTER_SC_APC"].sap, "2200005618")
        self.assertEqual(rows["ADAPTER_SC_APC"].qty_order, 98)
        self.assertEqual(rows["OSLONKA_SPAWU"].qty_order, 214)

    def test_engine_adds_hdpe_uv_pipe_accessories_for_dac_or_ground_microcable_at_oap(self):
        from ftth_bom.engine import run_analysis

        result = run_analysis(
            gpkg_path=JEDLNIA,
            catalog_path=CATALOG,
            bell_path=BELL,
            task_name="test_jedlnia_oap_hdpe_uv",
        )
        rows = {row.material_class: row for row in result.order_rows}

        self.assertEqual(result.project.metadata["oap_hdpe_uv_pipe_count"], 4)
        self.assertEqual(result.project.metadata["oap_hdpe_uv_energy_pipe_count"], 4)
        self.assertEqual(rows["HDPE_UV_40"].sap, "2200015460")
        self.assertEqual(rows["HDPE_UV_40"].qty_order, 20)
        self.assertEqual(rows["DYSTANS_HDPE_UV"].sap, "2200028483")
        self.assertEqual(rows["DYSTANS_HDPE_UV"].qty_order, 16)
        self.assertIn("91", rows["TASMA_STALOWA"].qty_pw)
        self.assertIn("121", rows["KLAMRA_TASMY"].qty_pw)

    def test_reader_counts_hdpe_uv_pipes_for_dac_at_oap_but_not_adss_service_drops(self):
        from ftth_bom.gpkg_reader import read_gpkg_project

        debe = read_gpkg_project(DEBE)
        dpw = read_gpkg_project(DPW_13171)

        self.assertEqual(debe.metadata["oap_hdpe_uv_pipe_count"], 0)
        self.assertEqual(dpw.metadata["oap_hdpe_uv_pipe_count"], 3)

    def test_engine_applies_dpw_client_feedback_rules(self):
        from ftth_bom.engine import run_analysis

        result = run_analysis(
            gpkg_path=DPW,
            catalog_path=CATALOG,
            bell_path=BELL,
            task_name="test_dpw_feedback",
        )
        rows = {row.material_class: row for row in result.order_rows}

        self.assertEqual(result.project.metadata["pigtail_splices"], 88)
        self.assertEqual(result.project.metadata["splice_sleeves"], 160)
        self.assertEqual(rows["DAC_2J"].qty_order, 9520)
        self.assertEqual(rows["MIKROKABEL_12J"].qty_order, 780)
        self.assertEqual(rows["MIKROKABEL_24J"].qty_order, 425)
        self.assertEqual(rows["MIKROKABEL_72J"].qty_order, 1150)
        self.assertEqual(rows["HDPE_40"].qty_order, 355)
        self.assertEqual(rows["MIKRORURKA_12_8"].qty_order, 2065)
        self.assertEqual(rows["ADAPTER_SC_APC"].sap, "2200005618")
        self.assertEqual(rows["ADAPTER_SC_APC"].qty_order, 88)
        self.assertEqual(rows["FUNDAMENT_PSB_H"].sap, "2200008667")
        self.assertEqual(rows["FUNDAMENT_PSB_H"].qty_order, 1)
        self.assertEqual(rows["NAKLEJKA_SI_50X72"].sap, "2200029709")
        self.assertEqual(rows["NAKLEJKA_SI_50X72"].qty_order, 16)
        self.assertEqual(rows["PIANKA_MD"].sap, "2200000503")
        self.assertEqual(rows["PIANKA_MD"].qty_order, 5)
        self.assertEqual(rows["ZAMEK_ABLOY"].sap, "2200029956")
        self.assertEqual(rows["ZAMEK_ABLOY"].qty_order, 7)
        self.assertEqual(rows["ZLACZKA_MIKRO_12"].sap, "2200008194")
        self.assertEqual(rows["ZLACZKA_MIKRO_12"].qty_order, 40)

    def test_reader_counts_existing_feeder_splices_from_real_okh_rows(self):
        from ftth_bom.gpkg_reader import read_gpkg_project

        project = read_gpkg_project(DPW_13171)

        self.assertEqual(project.metadata["pigtail_splices_raw"], 141)
        self.assertEqual(project.metadata["fiber_splice_connections"], 104)
        self.assertEqual(project.metadata["existing_feeder_splice_connections"], 2)
        self.assertEqual(project.metadata["splice_sleeves"], 245)

    def test_exporter_writes_xlsx_with_valid_sheet_names(self):
        from ftth_bom.engine import run_analysis
        from ftth_bom.exporters import write_xlsx

        result = run_analysis(
            gpkg_path=JEDLNIA,
            catalog_path=CATALOG,
            bell_path=BELL,
            task_name="test_export",
        )
        out = ROOT / "outputs" / "test_export.xlsx"
        write_xlsx(result, out)

        self.assertTrue(out.exists())

    def test_exporter_includes_project_layer_takeoff_in_pw_sheet(self):
        from openpyxl import load_workbook

        from ftth_bom.engine import run_analysis
        from ftth_bom.exporters import write_xlsx

        result = run_analysis(
            gpkg_path=JEDLNIA,
            catalog_path=CATALOG,
            bell_path=BELL,
            task_name="test_export_takeoff",
        )
        out = ROOT / "outputs" / "test_export_takeoff.xlsx"
        write_xlsx(result, out)

        wb = load_workbook(out, read_only=True, data_only=True)
        rows = ["\t".join("" if value is None else str(value) for value in row) for row in wb["Przedmiar PW"].iter_rows(values_only=True)]
        self.assertTrue(any("Skaner GPKG" in row and "MIKRORURKA_12_8" in row for row in rows))

    def test_web_upload_guard_does_not_bool_fieldstorage_objects(self):
        from web_app import has_upload_file

        class FieldLike:
            filename = "project.gpkg"

            def __bool__(self):
                raise TypeError("Cannot be converted to bool.")

        self.assertTrue(has_upload_file(FieldLike()))
        self.assertFalse(has_upload_file(None))

    def test_web_result_hides_material_preferences_panel(self):
        from ftth_bom.engine import run_analysis
        from web_app import render_result

        result = run_analysis(
            gpkg_path=JEDLNIA,
            catalog_path=CATALOG,
            bell_path=BELL,
            task_name="test_web_without_preferences_panel",
        )

        html = render_result(result)

        self.assertNotIn("Baza materiałów", html)
        self.assertNotIn("Baza preferencji", html)

    def test_web_download_headers_are_safe_for_polish_filenames(self):
        from web_app import attachment_header, output_href

        filename = "lista_materialow_PW_Jedlnia_OPP03_z_nr_wstęga.xlsx"

        header = attachment_header(filename)

        header.encode("latin-1")
        self.assertNotIn("ę", header)
        self.assertIn("filename*=", header)
        self.assertIn("wst%C4%99ga.xlsx", header)
        self.assertIn("wst%C4%99ga.xlsx", output_href(Path(filename)))

    def test_staged_gpkg_path_copies_to_temp_and_cleans_up(self):
        import tempfile

        from ftth_bom.staging import staged_gpkg_path

        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "projekt.gpkg"
            source.write_bytes(b"gpkg-bytes")

            with staged_gpkg_path(source) as staged:
                self.assertNotEqual(source.resolve(), staged.resolve())
                self.assertTrue(staged.exists())
                self.assertEqual(staged.read_bytes(), b"gpkg-bytes")

            self.assertFalse(staged.exists())

    def test_run_analysis_reports_progress_stages(self):
        from ftth_bom.engine import run_analysis

        events = []
        run_analysis(
            gpkg_path=DPW,
            catalog_path=CATALOG,
            bell_path=BELL,
            task_name="test_progress",
            local_copy=True,
            progress=events.append,
        )

        percents = [percent for percent, _message in events]
        messages = " ".join(message.lower() for _percent, message in events)
        self.assertGreaterEqual(max(percents), 90)
        self.assertEqual(percents, sorted(percents))
        self.assertIn("kopiowanie", messages)
        self.assertIn("gpkg", messages)

    def test_web_progress_page_polls_status_endpoint(self):
        from web_app import render_progress_page

        html = render_progress_page("abc123")

        self.assertIn("/status/abc123", html)
        self.assertIn("progress-bar", html)


if __name__ == "__main__":
    unittest.main()
