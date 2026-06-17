from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Side, Border
from openpyxl.utils import get_column_letter

from ftth_bom.models import AnalysisResult


ORDER_HEADERS = [
    "Lp",
    "Status",
    "Kategoria",
    "SAP",
    "Nazwa materiału",
    "JM",
    "Ilość wg PW",
    "Ilość do zamówienia",
    "Dostawca / pozycja katalogowa",
    "Podstawa doboru",
    "Uwagi",
]


def export_result(result: AnalysisResult, output_dir: str | Path = "outputs") -> AnalysisResult:
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    safe_task = "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in result.task_name).strip("_")
    xlsx_path = output_dir / f"lista_materialow_{safe_task}.xlsx"
    pdf_path = output_dir / f"lista_materialow_{safe_task}.pdf"
    write_xlsx(result, xlsx_path)
    try:
        write_pdf(result, pdf_path)
    except ModuleNotFoundError:
        pdf_path = None
    result.xlsx_path = xlsx_path
    result.pdf_path = pdf_path
    return result


def write_xlsx(result: AnalysisResult, path: str | Path) -> Path:
    path = Path(path)
    wb = Workbook()
    ws = wb.active
    ws.title = "Zamówienie"
    _write_order_sheet(ws, result)
    _write_issues_sheet(wb.create_sheet("Do wyjaśnienia"), result)
    _write_takeoff_sheet(wb.create_sheet("Przedmiar PW"), result)
    _write_used_poles_sheet(wb.create_sheet("Slupy uzyte"), result)
    _write_history_sheet(wb.create_sheet("Baza preferencji"), result)
    _write_sap_sheet(wb.create_sheet("SAP copy"), result)
    wb.save(path)
    return path


def write_pdf(result: AnalysisResult, path: str | Path) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    path = Path(path)
    regular, bold = _register_fonts()
    doc = SimpleDocTemplate(
        str(path),
        pagesize=landscape(A4),
        leftMargin=10 * mm,
        rightMargin=10 * mm,
        topMargin=10 * mm,
        bottomMargin=10 * mm,
    )
    styles = getSampleStyleSheet()
    styles["Normal"].fontName = regular
    styles["Title"].fontName = bold
    small = ParagraphStyle("small", parent=styles["Normal"], fontName=regular, fontSize=7, leading=8, alignment=TA_LEFT)
    story = [
        Paragraph(f"Lista materiałów FTTH - {result.task_name}", styles["Title"]),
        Paragraph(f"Data opracowania: {datetime.now().strftime('%Y-%m-%d %H:%M')}", styles["Normal"]),
        Paragraph(f"Źródło GPKG: {result.project.source_path.name}", styles["Normal"]),
        Spacer(1, 4 * mm),
        Paragraph("Punkty do wyjaśnienia", styles["Heading2"]),
    ]
    issues_data = [["Temat", "Opis", "Rekomendacja"]]
    for issue in result.issues[:20]:
        issues_data.append([Paragraph(issue.topic, small), Paragraph(issue.description, small), Paragraph(issue.recommendation, small)])
    story.append(_pdf_table(issues_data, [45 * mm, 115 * mm, 95 * mm]))
    story += [Spacer(1, 5 * mm), Paragraph("Tabela zamówienia", styles["Heading2"])]
    order_data = [ORDER_HEADERS[:8]]
    for row in result.order_rows[:80]:
        order_data.append(
            [
                row.lp,
                row.status,
                row.category,
                row.sap,
                Paragraph(row.name, small),
                row.unit,
                row.qty_pw,
                row.qty_order,
            ]
        )
    story.append(_pdf_table(order_data, [9 * mm, 25 * mm, 28 * mm, 25 * mm, 85 * mm, 12 * mm, 32 * mm, 28 * mm]))
    doc.build(story)
    return path


def _write_order_sheet(ws, result: AnalysisResult) -> None:
    ws.append(ORDER_HEADERS)
    for row in result.order_rows:
        ws.append(
            [
                row.lp,
                row.status,
                row.category,
                row.sap,
                row.name,
                row.unit,
                row.qty_pw,
                row.qty_order,
                row.supplier,
                row.basis,
                row.notes,
            ]
        )
    _style_table(ws, len(ORDER_HEADERS))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    warn_fill = PatternFill("solid", fgColor="FCE4D6")
    for cells in ws.iter_rows(min_row=2):
        if cells[1].value == "DO POTWIERDZENIA":
            for cell in cells:
                cell.fill = warn_fill


def _write_issues_sheet(ws, result: AnalysisResult) -> None:
    ws.append(["Temat", "Opis problemu", "Rekomendacja"])
    for issue in result.issues:
        ws.append([issue.topic, issue.description, issue.recommendation])
    _style_table(ws, 3)


def _write_takeoff_sheet(ws, result: AnalysisResult) -> None:
    ws.append(["Sekcja", "Klasa/typ", "Model/nazwa", "Ilość", "Trasa m", "Instalacyjna m", "Uwagi"])
    for cable in result.project.cable_summary:
        ws.append(["Kable", cable.material_class, cable.model, cable.sections, cable.route_m, cable.install_m, f"{cable.fibers or ''}J"])
    for device in result.project.devices:
        ws.append(["Urządzenia", device.material_class, device.model, device.count, "", "", device.node])
    for duct in result.project.ducts:
        ws.append([duct.group, "", duct.name, duct.count, duct.length_m, "", ""])
    for reserve in result.project.reserves:
        ws.append([reserve.group, "", reserve.name, reserve.count, reserve.length_m, "", ""])
    for item in result.project.material_takeoff:
        ws.append(
            [
                "Skaner GPKG",
                item.material_class,
                item.name,
                item.count,
                item.length_m,
                "",
                f"{item.source_layer}: {'; '.join(item.examples)}",
            ]
        )
    _style_table(ws, 7)


def _write_used_poles_sheet(ws, result: AnalysisResult) -> None:
    ws.append(["Warstwa", "Id slupa", "Model", "X", "Y", "Kable", "Klasy kabli", "Odleglosc m"])
    for pole in result.project.used_poles:
        ws.append(
            [
                pole.source_layer,
                pole.pole_id,
                pole.model,
                pole.x,
                pole.y,
                "; ".join(pole.matched_cables),
                "; ".join(pole.material_classes),
                pole.distance_m,
            ]
        )
    _style_table(ws, 8)


def _write_history_sheet(ws, result: AnalysisResult) -> None:
    ws.append(["Klasa", "SAP", "Nazwa", "JM", "Suma historyczna", "Wystąpienia", "Liczba F/X", "Przykłady"])
    for pref in result.preferences:
        ws.append(
            [
                pref.material_class,
                pref.sap,
                pref.name,
                pref.unit,
                pref.total_qty,
                pref.occurrences,
                pref.fx_count,
                "; ".join(pref.examples),
            ]
        )
    _style_table(ws, 8)


def _write_sap_sheet(ws, result: AnalysisResult) -> None:
    ws.append(["SAP", "Nazwa materiału", "JM", "Ilość do zamówienia", "Status"])
    for row in result.sap_rows:
        ws.append([row.sap, row.name, row.unit, row.qty_order, row.status])
    _style_table(ws, 5)


def _style_table(ws, cols: int) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="D9E2F3")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows():
        for cell in row:
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    widths = [10, 22, 24, 16, 45, 10, 20, 18, 32, 55, 45]
    for idx in range(1, cols + 1):
        ws.column_dimensions[get_column_letter(idx)].width = widths[idx - 1] if idx <= len(widths) else 24


def _register_fonts() -> tuple[str, str]:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    fonts = Path(r"C:\Windows\Fonts")
    regular = fonts / "arial.ttf"
    bold = fonts / "arialbd.ttf"
    if regular.exists():
        pdfmetrics.registerFont(TTFont("Arial", str(regular)))
    if bold.exists():
        pdfmetrics.registerFont(TTFont("Arial-Bold", str(bold)))
    return ("Arial" if regular.exists() else "Helvetica", "Arial-Bold" if bold.exists() else "Helvetica-Bold")


def _pdf_table(data, widths):
    from reportlab.lib import colors
    from reportlab.platypus import Table, TableStyle

    table = Table(data, colWidths=widths, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D9E2F3")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FBFD")]),
            ]
        )
    )
    return table
