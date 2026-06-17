import math
import sqlite3
import sys
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


BASE = Path(r"C:\Users\Pawel Z\Documents\Codex\2026-05-28\files-mentioned-by-the-user-historia")
INPUTS = BASE / "inputs"
OUTPUTS = BASE / "outputs"
GPKG = INPUTS / "PW Jedlnia OPP03 z nr wstęga.gpkg"
CATALOG = INPUTS / "KATALOGI PT_03_2026.xlsx"
HISTORY = INPUTS / "historia zamówień.xlsx"


def register_fonts():
    fonts = Path(r"C:\Windows\Fonts")
    regular = fonts / "arial.ttf"
    bold = fonts / "arialbd.ttf"
    if regular.exists():
        pdfmetrics.registerFont(TTFont("Arial", str(regular)))
    if bold.exists():
        pdfmetrics.registerFont(TTFont("Arial-Bold", str(bold)))
    return ("Arial" if regular.exists() else "Helvetica", "Arial-Bold" if bold.exists() else "Helvetica-Bold")


FONT, FONT_BOLD = register_fonts()


def fmt_sap(value):
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return ""
    return str(int(value)) if isinstance(value, float) else str(value)


def fmt_qty(value):
    if value is None or value == "":
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def catalog_map():
    cat = pd.read_excel(CATALOG, sheet_name="WYNIK")
    cat["Materiał SAP"] = pd.to_numeric(cat["Materiał SAP"], errors="coerce")

    preferences = {
        2200023905: "CORNING",  # project OAP is Corning and delivery is shortest
        2200023710: "BELOS",
        2200028041: "BELOS",
        2200006269: "OPTOMER",
        2200005618: "FCA",
        2200003990: "CORNING",
        2200003276: "ROCK",
        2200003277: "ROCK",
        2200003274: "ROCK",
        2200003402: "YOFC",
        2200003406: "YOFC",
        2200003404: "FCA",
        2200003264: "YOFC",
        2200003266: "YOFC",
        2200022783: "YOFC",
        2200008333: "OPTOMER",
        2200015460: "OPTOMER",
        2200004586: "FCA",
        2200004524: "FCA",
        2200003758: "LANDL",
        2200003756: "LANDL",
        2200003716: "LANDL",
        2200023869: "FCA",
        2200003817: "LANDL",
    }

    out = {}
    for sap, rows in cat.dropna(subset=["Materiał SAP"]).groupby("Materiał SAP"):
        sap_int = int(sap)
        selected = rows
        pref = preferences.get(sap_int)
        if pref:
            mask = rows[["Nazwa", "Opis poz Dostawcy dł", "Opis długi prd.", "Nr poz Dostawcy w PZ"]].fillna("").astype(str).agg(" ".join, axis=1).str.upper().str.contains(pref)
            if mask.any():
                selected = rows.loc[mask]
        row = selected.iloc[0]
        out[sap_int] = {
            "sap": sap_int,
            "opis_sap": row.get("Opis materiału SAP", ""),
            "nr_dostawcy": row.get("Nr poz Dostawcy w PZ", ""),
            "opis_dostawcy": row.get("Opis poz Dostawcy dł", ""),
            "opis_dlugi": row.get("Opis długi prd.", ""),
            "dostawca": row.get("Nazwa", ""),
            "kategoria": row.get("podkategoria 2", ""),
            "jm_katalog": row.get("Jn zakupowa PZ", ""),
            "termin": row.get("Termin dostawy - dni", ""),
        }
    return out


def history_opp03():
    hist = pd.read_excel(HISTORY, sheet_name="Sheet1")
    hist = hist[hist["NAZWA"].astype(str).str.upper().eq("OPP 03 JEDLNIA")].copy()
    hist["Ilosc_Zamówiona"] = pd.to_numeric(hist["Ilosc_Zamówiona"], errors="coerce").fillna(0)
    group = (
        hist.groupby(["NUMER \nINDEKSU\n(SAP)", "Nazwa Towaru", "JM"], dropna=False)
        .agg(ilosc=("Ilosc_Zamówiona", "sum"), wystapienia=("Nazwa Towaru", "size"))
        .reset_index()
        .sort_values("Nazwa Towaru")
    )
    return group


def project_metrics():
    con = sqlite3.connect(GPKG)
    con.row_factory = sqlite3.Row
    cur = con.cursor()
    metrics = {}

    metrics["cables"] = [dict(r) for r in cur.execute("""
        with per_cable as (
          select odcinek_kabla, model_kabla, liczba_wlokien,
                 max(coalesce(dl_instalacyjna,0)) dl_instalacyjna,
                 max(coalesce(dl_optyczna,0)) dl_optyczna,
                 max(coalesce(dl_trasowa,0)) dl_trasowa
          from [Kable Światłowodowe]
          group by odcinek_kabla, model_kabla, liczba_wlokien
        )
        select model_kabla, liczba_wlokien, count(*) liczba_odcinkow,
               round(sum(dl_trasowa), 2) trasa_m,
               round(sum(dl_optyczna), 2) optyczna_m,
               round(sum(dl_instalacyjna), 2) instalacyjna_m
        from per_cable
        group by model_kabla, liczba_wlokien
        order by liczba_wlokien, model_kabla
    """)]

    metrics["devices"] = [dict(r) for r in cur.execute("""
        select typ_elementu, producent, model_urzadzenia, typ_obiektu, count(*) ilosc
        from [Urządzenia Pasywne]
        group by typ_elementu, producent, model_urzadzenia, typ_obiektu
        order by typ_elementu, model_urzadzenia
    """)]

    metrics["ducts"] = [dict(r) for r in cur.execute("""
        select typ_elementu, producent, oznaczenie, count(*) odcinki, round(sum(dlugosc),2) dlugosc_m
        from [Odcinki Kanalizacji]
        group by typ_elementu, producent, oznaczenie
        order by typ_elementu, oznaczenie
    """)]

    metrics["mufa_osd0136_cables"] = [dict(r) for r in cur.execute("""
        select odcinek_kabla, model_kabla, liczba_wlokien, od, do,
               max(coalesce(dl_instalacyjna,0)) dl_instalacyjna
        from [Kable Światłowodowe]
        where od like '%OSD0136%' or do like '%OSD0136%'
        group by odcinek_kabla, model_kabla, liczba_wlokien, od, do
        order by model_kabla, odcinek_kabla
    """)]

    con.close()
    return metrics


def row(lp, status, category, sap, name, jm, qty_pw, qty_order, basis, notes="", catalog=None):
    c = catalog.get(sap, {}) if catalog and isinstance(sap, int) else {}
    return {
        "Lp": lp,
        "Status": status,
        "Kategoria": category,
        "SAP": fmt_sap(sap),
        "Nazwa materiału": name or c.get("opis_sap", ""),
        "JM": jm,
        "Ilość wg PW": qty_pw,
        "Ilość do zamówienia": qty_order,
        "Dostawca / pozycja katalogowa": f"{c.get('dostawca','')} | {c.get('nr_dostawcy','')}".strip(" |"),
        "Podstawa doboru": basis,
        "Uwagi": notes,
    }


def build_rows(catalog):
    rows = []
    add = lambda *args, **kwargs: rows.append(row(len(rows) + 1, *args, catalog=catalog, **kwargs))

    add("ZAMÓWIĆ", "Kable", 2200003264, "KAB.ŚWIATŁ.DAC 2J", "m", "13 125", "13 200",
        "PW: 83 unikalne odcinki DAC 2J; historia OPP03: 13 200 m.",
        "Zaokrąglenie jak w historii; kabel DAC Ø ok. 5,9 mm.")
    add("DO POTWIERDZENIA", "Kable", 2200003266, "KAB.ŚWIATŁ.DAC 6J zamiast DAC 4J", "m", "60 m DAC 4J", "60",
        "PW ma DAC 4J, ale w katalogu brak DAC 4J; historia OPP03 używa DAC 6J 60 m.",
        "Wymaga akceptacji zamiennika przez projektanta/inwestora.")
    add("DO POTWIERDZENIA", "Kable", 2200022783, "KAB.ABON.ADSS.DROP FLEXTUBE 1x2J 657A1", "m", "360 trasy / 0 instal.", "400",
        "PW: 11 planowanych kabli napowietrznych 2J; długość instalacyjna w PW = 0.",
        "Prawdopodobny błąd PW. Przyjęto trasę 360 m i zaokrąglono do 400 m.")
    add("ZAMÓWIĆ", "Kable", 2200003276, "KAB.ŚWIATŁ.ADSS 12J (3x4J) 657A1 3kN", "m", "645", "650",
        "PW: 2 odcinki ADSS 12J, historia OPP03: 650 m.",
        "Kabel Ø ok. 10,8 mm; dobór uchwytów 8-12 mm.")
    add("ZAMÓWIĆ", "Kable", 2200003277, "KAB.ŚWIATŁ.ADSS 24J (6x4J) 657A1 3kN", "m", "300", "310",
        "PW: 2 odcinki ADSS 24J, historia OPP03: 310 m.",
        "Kabel Ø ok. 10,8 mm; dobór uchwytów 8-12 mm.")
    add("ZAMÓWIĆ", "Kable", 2200003274, "KAB.ŚWIATŁ.ADSS 36J (6x6J) 657A1 3kN", "m", "400", "410",
        "PW: 3 odcinki ADSS 36J, historia OPP03: 410 m.",
        "Kabel Ø ok. 10,8 mm; dobór uchwytów 8-12 mm.")
    add("ZAMÓWIĆ", "Kable", 2200003402, "MIKROKABEL 12J-1x12", "m", "610", "620",
        "PW: 2 odcinki MI-MKF 12J; historia OPP03: 620 m.",
        "G.652D, mikrokabel Ø ok. 5,4 mm.")
    add("DO POTWIERDZENIA", "Kable", 2200003406, "MIKROKABEL 24J-2x12", "m", "325", "330",
        "PW: 1 odcinek MI-MKF 24J; brak tej pozycji w historii OPP03.",
        "Potwierdzić, czy odcinek OSD0127-ALEKSANDJMAZ/OSD0001 jest w zakresie tego zamówienia.")
    add("DO POTWIERDZENIA", "Kable", 2200003404, "MIKROKABEL 144J-12x12", "m", "545", "550",
        "PW: 3 odcinki MI-MKF 144J; historia OPP03 pokazuje 1300 m.",
        "Duża różnica PW vs historia. Nie podbijam do 1300 m bez potwierdzenia.")

    add("ZAMÓWIĆ", "Kanalizacja", 2200004586, "MIKRORURKA 12/8", "m", "1 155", "1 160",
        "PW: FP-MR-G-12/8, rurociąg 630 m + kanalizacja wtórna 525 m.",
        "Historia OPP03 ma 700 m, ale PW wskazuje większą długość.")
    add("ZAMÓWIĆ", "Kanalizacja", 2200004524, "MIKRORURKA 14/10", "m", "308", "500",
        "PW: FP-MR-G-14/10 308 m; historia OPP03: 500 m.",
        "Przyjęto historyczne zaokrąglenie/zapewnienie zapasu.")
    add("ZAMÓWIĆ", "Kanalizacja", 2200008333, "RURA HDPEp FI40x3,7 CZARNA", "m", "987", "1 250",
        "PW: jawne odcinki RHDPEp 40x3,7 = 917 m + rura osłonowa 70 m; historia OPP03: 1250 m.",
        "Nie wliczam generycznych odcinków doziemnych jako rury.")
    add("ZAMÓWIĆ", "Kanalizacja", 2200015460, "RURA HDPE-UV FI 40/3,7", "m", "brak osobnej warstwy", "35",
        "Historia OPP03 i praktyka ochrony podejść na słupach.",
        "Pozycja akcesoryjna dla odcinków narażonych na UV.")
    add("ZAMÓWIĆ", "Kanalizacja", 2200004707, "TAŚMA OST.TO-TKT/10 ~UWAGA KABEL~", "m", "~2 015 wykopów/przepustów", "3 000",
        "Historia OPP03: 3000 m; PW: suma jawnych odcinków ziemnych ok. 2,0 km.",
        "Zapas na dojścia, korekty trasy i odtworzenia.")
    add("ZAMÓWIĆ", "Kanalizacja", 2200008194, "ZŁĄCZKA PROSTA 12MM MIKRORUREK", "szt", "19 otworów 12/8", "10",
        "Historia OPP03: 10 szt.",
        "Jeśli wykonawca nie ma magazynu, warto podnieść do 20 szt.")
    add("ZAMÓWIĆ", "Kanalizacja", 2200004787, "ZŁĄCZKA PROSTA 14MM", "szt", "2 otwory 14/10", "25",
        "Historia OPP03: 25 szt.",
        "Wysoka liczba historyczna, prawdopodobnie pakiet/minimum lub zapas montażowy.")
    add("ZAMÓWIĆ", "Kanalizacja", 2200008140, "ZŁĄCZKA DO RUR PCV ZRS-40", "szt", "RHDPE/PCV - do przejść", "6",
        "Historia OPP03: 6 szt.",
        "")
    add("ZAMÓWIĆ", "Kanalizacja", 2200005354, "USZCZ.MIKROR.12MM FP-UMD-12/6,5-8,0", "szt", "12/8", "8",
        "Historia OPP03: 8 szt.; zakres pasuje do mikrorur 12/8.",
        "Dla 14/10 brak pozycji w historii OPP03 - do sprawdzenia na etapie realizacji.")
    add("DO POTWIERDZENIA", "Kanalizacja", 2200005351, "USZCZELN.DZIEL.MIKROR.UMD-14/6,5-8,0MM", "szt", "2 odcinki FP-MR-G-14/10, 308 m", "4",
        "PW: mikrorurka 14/10 występuje na 2 odcinkach; historia OPP03 nie miała tej pozycji, co było powodem pierwotnego pominięcia.",
        "Zakres 6,5-8,0 mm pasuje do mikrokabla 144J Ø ok. 7,8 mm. Jeśli w 14/10 idzie kabel Ø 5,0-6,5 mm, właściwszy będzie SAP 2200030172.")

    add("ZAMÓWIĆ", "Punkty pasywne", 2200028376, "MUFA ŚWIATŁ.SSC2110-FM_48", "szt", "1", "1",
        "PW: 1 mufa OSD0136; w historii OPP03 też 1 szt.",
        "Dobór uszczelnień: w OSD0136 są 2x MI-MKF 144J Ø ok. 7,8 mm i 9x DAC 2J Ø ok. 5,9 mm; zestaw SSC2110-FM_48 zawiera uszczelnienia 4-16, 5-13 i 12x drop 3-7, więc wystarcza bez dodatkowych uszczelnień.")
    add("ZAMÓWIĆ", "Punkty pasywne", 2200023135, "PRZEŁ.ŚWIATŁ.PSB-H-144-GM-K144-00", "szt", "1", "1",
        "PW: OPP0013 jako PSB-H-144-GM; historia OPP03: 1 szt.",
        "")
    add("ZAMÓWIĆ", "Punkty pasywne", 2200008667, "FUNDAMENT PREFABRYKOWANY PSB-H", "szt", "1", "1",
        "Do szafy PSB-H; historia OPP03: 1 szt.",
        "")
    add("ZAMÓWIĆ", "Punkty pasywne", 2200023700, "SŁUPEK AB.SUS-PH BEZ KOMUTACJI 96/48SP", "szt", "4", "4",
        "PW: 4 przełącznice SUS-PH-S na słupkach; historia OPP03: 4 szt.",
        "")
    add("ZAMÓWIĆ", "Punkty pasywne", 2200023905, "SKRZYNKA OAP-8", "szt", "7", "7",
        "PW: 7 OAP_8 na słupach; historia OPP03: 7 szt.",
        "Wybrałem wariant Corning, zgodny z producentem w PW i najkrótszym terminem w katalogu.")
    add("ZAMÓWIĆ", "Punkty pasywne", 2200023710, "STELAŻ ZAPASU KAB.OAP 8/48", "szt", "7", "7",
        "PW: zapasy projektowane przy OAP; historia OPP03: 7 szt.",
        "")
    add("ZAMÓWIĆ", "Punkty pasywne", 2200028041, "DYSTANS 150MM STELAŻA ZAPASU KAB OAP", "szt", "7 OAP", "12",
        "Historia OPP03: 12 szt.",
        "Zapas na montaż stelaży na słupach energetycznych.")
    add("ZAMÓWIĆ", "Punkty pasywne", 2200006438, "STELAŻ ZAPASU KAB.STZK-60", "szt", "1", "1",
        "Historia OPP03: 1 szt.; zapas przy punkcie/mufie.",
        "")
    add("ZAMÓWIĆ", "Optyka", 2200006269, "SPLITTER PLC BLACKBOX 1X64/657/2.0/SCA", "szt", "1", "1",
        "PW: 1 spliter SPL1x64/1216/SCA w OPP0013; historia OPP03: 1 szt.",
        "")
    add("ZAMÓWIĆ", "Optyka", 2200005618, "ADAPTER SC/APC", "szt", "porty PSB/komutacja", "120",
        "Historia OPP03: 120 szt.; zgodne z używanym standardem SC/APC.",
        "OAP-8 i PSB-H mają część adapterów w wyposażeniu, ale historia wskazuje dodatkową pulę.")
    add("ZAMÓWIĆ", "Optyka", 2200005414, "PIGTAIL 657 SC/APC DŁ.3M", "szt", "spawy/pola komutacyjne", "163",
        "Historia OPP03: 163 szt.",
        "")
    add("ZAMÓWIĆ", "Optyka", 2200003990, "OSŁONKA SPAWU ŚWIATŁOWODU OS-45", "szt", "spawy", "280",
        "Historia OPP03: 280 szt.; liczba spawów wynika z OPP, OSD i kabli abonenckich.",
        "")
    add("ZAMÓWIĆ", "Optyka", 2200005195, "KAPTUREK TERMOKURCZLIWY NA KAB.DAC", "szt", "84 kable DAC", "86",
        "PW: 83x DAC 2J + 1x DAC 4J; historia OPP03: 86 szt.",
        "")

    add("ZAMÓWIĆ", "Napowietrzne", 2200003817, "UCHWYT ODCIĄGOWY UOO/10-W (8-12mm)", "szt", "ADSS 12/24/36", "50",
        "Historia OPP03: 50 szt.; pasuje do ADSS 12/24/36 o średnicy ok. 10,8 mm.",
        "Ilość historyczna zawiera zapas i mocowania pośrednie, nie tylko końce odcinków.")
    add("DO POTWIERDZENIA", "Napowietrzne", 2200023869, "UCHWYT ODCIĄGOWY KLINOWY FCA UOZ 5-7mm", "szt", "11 dropów ADSS 2J", "24",
        "ADSS drop 2J ma średnicę ok. 5,8 mm; przyjęto 2 uchwyty na kabel + zapas.",
        "Zamawiać razem z ADSS 2J po potwierdzeniu błędu długości w PW.")
    add("OPCJONALNIE", "Napowietrzne", 2200003716, "HAK UNIWERSALNY", "szt", "mocowania na słupach", "24",
        "Akcesorium dla uchwytów odciągowych, jeżeli nie wykorzystuje się istniejących haków/wsporników.",
        "Nie było w historii OPP03, więc oznaczam jako opcjonalne.")
    add("OPCJONALNIE", "Napowietrzne", 2200003758, "TAŚMA STALOWA TSM/20-07-J", "rol", "mocowania słupowe", "1",
        "Dla mocowania osprzętu na słupach taśmą 20 mm.",
        "Nie było w historii OPP03; zamawiać, jeśli brak na magazynie.")
    add("OPCJONALNIE", "Napowietrzne", 2200003756, "KLAMRA DO TAŚMY STALOWEJ TSK/20-J", "pak", "taśma stalowa 20 mm", "1",
        "Komplet z taśmą stalową.",
        "Nie było w historii OPP03; zamawiać, jeśli brak na magazynie.")
    add("ZAMÓWIĆ", "Napowietrzne", 2200028483, "UCHWYT DYSTANS.RUR 25-50 NA SŁ.ŻELBET", "szt", "podejścia rur na słupach", "24",
        "Historia OPP03: 24 szt.",
        "")

    add("ZAMÓWIĆ", "Drobne/instalacyjne", 2200029956, "ZAMEK ABLOY CL704B_S-I_Z01.01.00_CENTRUM", "szt", "1 PSB + 4 SUS", "5",
        "Historia OPP03: 5 szt.; liczba zgodna z 5 obudowami wymagającymi zamka.",
        "")
    add("ZAMÓWIĆ", "Drobne/instalacyjne", 2200004672, "TAŚMA IZOLACYJNA SCOTCH 88T 19X11", "szt", "montaż", "5",
        "Historia OPP03: 5 szt.",
        "")
    add("ZAMÓWIĆ", "Drobne/instalacyjne", 2200000503, "PIANKA USZCZELNIAJĄCA MD+ 310ml", "szt", "uszczelnienia przepustów", "5",
        "Historia OPP03: 5 szt.",
        "")
    add("ZAMÓWIĆ", "Drobne/instalacyjne", 2200005208, "ŁĄCZNIK JEDNOŻYŁOWY PRZELOTOWY UY 2", "szt", "drobne połączenia", "20",
        "Historia OPP03: 20 szt.",
        "")
    add("ZAMÓWIĆ", "Drobne/instalacyjne", 2200006266, "OSŁ.KAB.KM2-PL DO 10-PAR ŻELOWANA", "szt", "osłony kablowe", "2",
        "Historia OPP03: 2 szt.",
        "")

    return rows


def apply_domain_rules(rows, catalog):
    # Korekty wynikajace z zasad wykonawczych przekazanych po pierwszej wersji BOM.
    rows = [r for r in rows if r["SAP"] not in {"2200022783", "2200023869"}]

    def update(sap, **fields):
        sap = str(sap)
        for r in rows:
            if r["SAP"] == sap:
                r.update(fields)
                return r
        raise KeyError(f"Missing SAP {sap}")

    update(
        2200003266,
        Status="ZAMÓWIĆ",
        **{
            "Ilość wg PW": "60 m DAC 4J",
            "Ilość do zamówienia": "60",
            "Podstawa doboru": "PW ma 60 m DAC 4J; zgodnie z zasadą wykonawczą każdy DAC 4J zamieniam na DAC 6J.",
            "Uwagi": "Katalog nie zawiera DAC 4J.",
        },
    )
    update(
        2200015460,
        **{
            "Ilość wg PW": "7 słupów z wyjściem HDPE",
            "Ilość do zamówienia": "35",
            "Podstawa doboru": "7 słupów z OAP/wyjściem rury na słup x 5 m HDPE-UV.",
            "Uwagi": "Odcinek ochronny UV przy wyjściu z ziemi lub studni na słup.",
        },
    )
    update(
        2200008194,
        **{
            "Ilość wg PW": "1 155 m / 50 m",
            "Ilość do zamówienia": "24",
            "Podstawa doboru": "Zasada wykonawcza: 1 złączka prosta 12 mm na każde rozpoczęte 50 m mikrorurki 12/8; ceil(1155/50)=24.",
            "Uwagi": "Zastępuje wcześniejsze przyjęcie z historii OPP03.",
        },
    )
    update(
        2200004787,
        **{
            "Ilość wg PW": "308 m / 50 m",
            "Ilość do zamówienia": "7",
            "Podstawa doboru": "Zasada wykonawcza: 1 złączka prosta 14 mm na każde rozpoczęte 50 m mikrorurki 14/10; ceil(308/50)=7.",
            "Uwagi": "Zastępuje wcześniejsze historyczne 25 szt.",
        },
    )
    update(
        2200008140,
        **{
            "Ilość wg PW": "7 wyjść HDPE na słup",
            "Ilość do zamówienia": "7",
            "Podstawa doboru": "Jedna złączka rury HDPE 40 mm dla każdego wyjścia HDPE z ziemi/studni na słup.",
            "Uwagi": "Powiązane z pozycją HDPE-UV 5 m/słup.",
        },
    )
    update(
        2200005354,
        **{
            "Ilość wg PW": "7 OAP + 1 SSC",
            "Ilość do zamówienia": "8",
            "Podstawa doboru": "Uszczelnienia mikrorurek stosowane tylko przy OAP/BPEO/FIST/SSC; w tym projekcie 7 OAP i 1 SSC dla torów 12/8.",
            "Uwagi": "Wymagane uszczelnienie dwudzielne.",
        },
    )
    update(
        2200005351,
        Status="ZAMÓWIĆ",
        **{
            "Ilość wg PW": "1 wejście 14/10 przy SSC",
            "Ilość do zamówienia": "1",
            "Podstawa doboru": "PW ma 308 m mikrorurki 14/10; uszczelnienie dwudzielne 14 mm przyjmuję tylko przy mufie/SSC, a nie na każdy odcinek trasy.",
            "Uwagi": "Zakres 6,5-8,0 mm pasuje do mikrokabla 144J ok. 7,8 mm. Dla kabla 5,0-6,5 mm stosować SAP 2200030172.",
        },
    )
    update(
        2200005618,
        **{
            "Ilość wg PW": "OPP: 96 HH + 2 dosył",
            "Ilość do zamówienia": "98",
            "Podstawa doboru": "W OPP stosujemy adaptery; ilość z projektu = 96 HH + 2 włókna dosyłowe (1 włókno na rozpoczęte 64 HH).",
            "Uwagi": "OAP dostajemy wyposażone w adaptery, więc nie doliczam adapterów OAP.",
        },
    )
    update(
        2200005414,
        **{
            "Ilość wg PW": "98 OPP + 19 OAP",
            "Ilość do zamówienia": "117",
            "Podstawa doboru": "Z warstwy włókien: 98 zakończeń pigtailowych w OPP i 19 w OAP. SUS-PH bez komutacji liczę jako włókno-włókno, bez pigtaili.",
            "Uwagi": "OAP ma adaptery fabrycznie, ale pigtail należy domówić.",
        },
    )
    update(
        2200003990,
        **{
            "Ilość wg PW": "spawy OPP/OAP/SUS/SSC",
            "Ilość do zamówienia": "280",
            "Podstawa doboru": "Historia OPP03: 280 szt.; po korekcie pigtaili nadal zostają spawy włókno-włókno w SUS-PH bez komutacji i w SSC.",
            "Uwagi": "Nie obniżam tej pozycji, bo osłonki obejmują także spawy poza pigtailami.",
        },
    )
    update(
        2200005195,
        **{
            "Ilość wg PW": "83 DAC 2J + 1 DAC6J",
            "Ilość do zamówienia": "84",
            "Podstawa doboru": "Kapturek termokurczliwy na DAC: 1 szt. na kabel DAC, tylko z jednej strony.",
            "Uwagi": "DAC 4J z PW zastąpiony DAC 6J.",
        },
    )
    update(
        2200003817,
        **{
            "Ilość wg PW": "25 słupów x 2 haki",
            "Ilość do zamówienia": "50",
            "Podstawa doboru": "Zamawiamy tyle uchwytów odciągowych, ile haków; 25 słupów x 2 = 50.",
            "Uwagi": "UOO/10-W dla kabli ADSS 12/24/36 o średnicy ok. 10,8 mm.",
        },
    )
    update(
        2200003716,
        Status="ZAMÓWIĆ",
        **{
            "Ilość wg PW": "25 słupów x 2",
            "Ilość do zamówienia": "50",
            "Podstawa doboru": "Zasada wykonawcza: na jeden słup zamawiamy dwa haki.",
            "Uwagi": "Liczba haków równa liczbie uchwytów odciągowych.",
        },
    )

    simplex = row(
        0,
        "ZAMÓWIĆ",
        "Kanalizacja",
        2200005322,
        "USZCZELKA JM SIMPLEX FI40 9-14",
        "szt",
        "7 wyjść HDPE z jedną mikrorurką",
        "7",
        "Uszczelnienie pomiędzy HDPE 40 a jedną mikrorurką: simplex; 1 szt. na każde wyjście HDPE na słup.",
        "Dla dwóch lub trzech mikrorurek stosować inny typ uszczelnienia.",
        catalog,
    )
    insert_at = next((i + 1 for i, r in enumerate(rows) if r["SAP"] == "2200005351"), len(rows))
    rows.insert(insert_at, simplex)

    for idx, r in enumerate(rows, start=1):
        r["Lp"] = idx
    return rows


def updated_issues():
    return [
        {
            "Temat": "Przyłącza napowietrzne ADSS 2J",
            "Opis": "PW pokazuje 11 planowanych przyłączy napowietrznych ADSS 2J o trasie 360 m, ale na zadaniach FTTH SI nie zostawiamy zapasów abonenckich kabli napowietrznych.",
            "Rekomendacja": "Nie ujmuję w zamówieniu ADSS 2J ani uchwytów UOZ 5-7 mm dla tych przyłączy.",
        },
        {
            "Temat": "Adaptery i pigtaile",
            "Opis": "Z projektu odczytałem 96 HH, 2 włókna dosyłowe w OPP oraz 19 zakończeń pigtailowych w OAP. OAP ma adaptery fabrycznie; SUS-PH bez komutacji spawany jest włókno-włókno.",
            "Rekomendacja": "W BOM przyjmuję 98 adapterów SC/APC do OPP oraz 117 pigtaili SC/APC łącznie dla OPP i OAP.",
        },
        {
            "Temat": "MI-MKF 144J",
            "Opis": "PW daje 545 m mikrokabla 144J, natomiast historia OPP03 ma 1300 m.",
            "Rekomendacja": "Nie zwiększać automatycznie do 1300 m. Potwierdzić, czy historia obejmowała dodatkowy zakres albo czy aktualny GPKG jest po zmianach.",
        },
        {
            "Temat": "MI-MKF 24J",
            "Opis": "PW zawiera 325 m mikrokabla 24J na odcinku OSD0127-ALEKSANDJMAZ/OSD0001, ale historia OPP03 go nie zawiera.",
            "Rekomendacja": "Potwierdzić, czy ten odcinek należy do zakresu zamówienia OPP03.",
        },
        {
            "Temat": "Warstwy K_*",
            "Opis": "Warstwy koncepcyjne K Linia Napowietrzna/K Słup zawierają ok. 37,4 km linii i 1043 słupy, czyli ewidentnie szerszy zakres niż PW OPP03.",
            "Rekomendacja": "Nie wliczałem ich do BOM. Gdyby to miało być zamówienie dla całej koncepcji, trzeba zrobić oddzielny przedmiar.",
        },
        {
            "Temat": "Zestawienie Czynności i Materiałów",
            "Opis": "Tabela gotowego zestawienia w GPKG ma 0 rekordów.",
            "Rekomendacja": "BOM został policzony z warstw projektu i zweryfikowany historią zamówień.",
        },
    ]


ISSUES = [
    {
        "Temat": "ADSS 2J w PW",
        "Opis": "W warstwie kabli jest 11 planowanych kabli napowietrznych CTC ADSS 2J o sumie trasy 360 m, ale długość instalacyjna i optyczna wynoszą 0 m.",
        "Rekomendacja": "Potwierdzić z autorem PW. Do zamówienia wpisałem 400 m ADSS drop 2J i 24 uchwyty UOZ 5-7 mm jako pozycje do potwierdzenia.",
    },
    {
        "Temat": "DAC 4J",
        "Opis": "PW zawiera 60 m kabla DAC 4J, a katalog inwestora nie ma zatwierdzonej pozycji DAC 4J.",
        "Rekomendacja": "Historia OPP03 pokazuje 60 m DAC 6J. W tabeli wpisałem DAC 6J jako zamiennik do akceptacji.",
    },
    {
        "Temat": "MI-MKF 144J",
        "Opis": "PW daje 545 m mikrokabla 144J, natomiast historia OPP03 ma 1300 m.",
        "Rekomendacja": "Nie zwiększać automatycznie do 1300 m. Potwierdzić, czy historia obejmowała dodatkowy zakres albo czy aktualny GPKG jest po zmianach.",
    },
    {
        "Temat": "MI-MKF 24J",
        "Opis": "PW zawiera 325 m mikrokabla 24J na odcinku OSD0127-ALEKSANDJMAZ/OSD0001, ale historia OPP03 go nie zawiera.",
        "Rekomendacja": "Potwierdzić, czy ten odcinek należy do zakresu zamówienia OPP03.",
    },
    {
        "Temat": "Uszczelnienie mikrorurki 14 mm",
        "Opis": "PW ma 2 odcinki mikrorurki 14/10 o łącznej długości 308 m. Historia OPP03 zawiera mikrorurkę 14/10 i złączki 14 mm, ale nie zawiera uszczelnienia 14 mm.",
        "Rekomendacja": "Dodałem 4 szt. UMD-14/6,5-8,0 mm jako do potwierdzenia. Jeżeli w 14/10 będzie kabel 5,0-6,5 mm, trzeba zamienić na SAP 2200030172.",
    },
    {
        "Temat": "Warstwy K_*",
        "Opis": "Warstwy koncepcyjne K Linia Napowietrzna/K Słup zawierają ok. 37,4 km linii i 1043 słupy, czyli ewidentnie szerszy zakres niż PW OPP03.",
        "Rekomendacja": "Nie wliczałem ich do BOM. Gdyby to miało być zamówienie dla całej koncepcji, trzeba zrobić oddzielny przedmiar.",
    },
    {
        "Temat": "Zestawienie Czynności i Materiałów",
        "Opis": "Tabela gotowego zestawienia w GPKG ma 0 rekordów.",
        "Rekomendacja": "BOM został policzony z warstw projektu i zweryfikowany historią zamówień.",
    },
]


def write_xlsx(rows, issues, metrics, history):
    path = OUTPUTS / "lista_materialow_Jedlnia_OPP03.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Zamówienie"
    headers = list(rows[0].keys())
    ws.append(headers)
    for r in rows:
        ws.append([r[h] for h in headers])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_cells in ws.iter_rows(min_row=2):
        status = row_cells[1].value
        fill = None
        if status == "DO POTWIERDZENIA":
            fill = PatternFill("solid", fgColor="FFF2CC")
        elif status == "OPCJONALNIE":
            fill = PatternFill("solid", fgColor="E2F0D9")
        for cell in row_cells:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(top=thin, left=thin, right=thin, bottom=thin)
            if fill:
                cell.fill = fill
    widths = [5, 17, 18, 12, 34, 8, 16, 18, 38, 48, 58]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wi = wb.create_sheet("Do wyjaśnienia")
    wi.append(["Temat", "Opis", "Rekomendacja"])
    for issue in issues:
        wi.append([issue["Temat"], issue["Opis"], issue["Rekomendacja"]])
    for cell in wi[1]:
        cell.fill = PatternFill("solid", fgColor="C65911")
        cell.font = header_font
    for col, width in enumerate([28, 80, 80], start=1):
        wi.column_dimensions[get_column_letter(col)].width = width
    for row_cells in wi.iter_rows():
        for cell in row_cells:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    wm = wb.create_sheet("Przedmiar PW")
    wm.append(["Sekcja", "Dane"])
    for cable in metrics["cables"]:
        wm.append(["Kable", str(cable)])
    for duct in metrics["ducts"]:
        wm.append(["Kanalizacja", str(duct)])
    for dev in metrics["devices"]:
        wm.append(["Urządzenia", str(dev)])
    for cable in metrics["mufa_osd0136_cables"]:
        wm.append(["Mufa OSD0136", str(cable)])
    wm.column_dimensions["A"].width = 18
    wm.column_dimensions["B"].width = 140
    for row_cells in wm.iter_rows():
        for cell in row_cells:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    wh = wb.create_sheet("Historia OPP03")
    wh.append(list(history.columns))
    for _, record in history.iterrows():
        wh.append([record[c] for c in history.columns])
    for col in range(1, len(history.columns) + 1):
        wh.column_dimensions[get_column_letter(col)].width = 24
    for cell in wh[1]:
        cell.fill = PatternFill("solid", fgColor="70AD47")
        cell.font = header_font

    wb.save(path)
    return path


def para(text, style):
    return Paragraph(str(text).replace("\n", "<br/>"), style)


def write_pdf(rows, issues):
    path = OUTPUTS / "lista_materialow_Jedlnia_OPP03.pdf"
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="BodyPL", fontName=FONT, fontSize=7.2, leading=8.4, alignment=TA_LEFT))
    styles.add(ParagraphStyle(name="SmallPL", fontName=FONT, fontSize=6.2, leading=7.2, alignment=TA_LEFT))
    styles.add(ParagraphStyle(name="TitlePL", fontName=FONT_BOLD, fontSize=15, leading=18, spaceAfter=6))
    styles.add(ParagraphStyle(name="HeadingPL", fontName=FONT_BOLD, fontSize=10, leading=12, spaceBefore=8, spaceAfter=4))

    doc = SimpleDocTemplate(
        str(path),
        pagesize=landscape(A4),
        rightMargin=8 * mm,
        leftMargin=8 * mm,
        topMargin=8 * mm,
        bottomMargin=8 * mm,
    )
    story = []
    story.append(Paragraph("Lista materiałów do zamówienia - PW Jedlnia OPP03", styles["TitlePL"]))
    story.append(Paragraph(f"Data opracowania: {datetime.now().strftime('%Y-%m-%d %H:%M')}. Źródła: GPKG PW, katalog PT_03_2026, historia zamówień OPP 03 Jedlnia.", styles["BodyPL"]))
    story.append(Spacer(1, 4))

    story.append(Paragraph("Punkty do wyjaśnienia", styles["HeadingPL"]))
    issue_data = [[para("Temat", styles["SmallPL"]), para("Opis", styles["SmallPL"]), para("Rekomendacja", styles["SmallPL"])]]
    for issue in issues:
        issue_data.append([para(issue["Temat"], styles["SmallPL"]), para(issue["Opis"], styles["SmallPL"]), para(issue["Rekomendacja"], styles["SmallPL"])])
    issue_table = Table(issue_data, colWidths=[38 * mm, 120 * mm, 120 * mm], repeatRows=1)
    issue_table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, 0), FONT_BOLD),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#C65911")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D9E2F3")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#FFF7EA")]),
    ]))
    story.append(issue_table)
    story.append(PageBreak())

    story.append(Paragraph("Tabela zamówienia", styles["HeadingPL"]))
    headers = ["Lp", "Status", "Kategoria", "SAP", "Nazwa materiału", "JM", "Ilość do zamówienia", "Podstawa doboru", "Uwagi"]
    data = [[para(h, styles["SmallPL"]) for h in headers]]
    for r in rows:
        data.append([
            para(r["Lp"], styles["SmallPL"]),
            para(r["Status"], styles["SmallPL"]),
            para(r["Kategoria"], styles["SmallPL"]),
            para(r["SAP"], styles["SmallPL"]),
            para(r["Nazwa materiału"], styles["SmallPL"]),
            para(r["JM"], styles["SmallPL"]),
            para(r["Ilość do zamówienia"], styles["SmallPL"]),
            para(r["Podstawa doboru"], styles["SmallPL"]),
            para(r["Uwagi"], styles["SmallPL"]),
        ])
    table = Table(data, colWidths=[8 * mm, 24 * mm, 23 * mm, 20 * mm, 47 * mm, 10 * mm, 22 * mm, 70 * mm, 54 * mm], repeatRows=1)
    table_style = [
        ("FONTNAME", (0, 0), (-1, 0), FONT_BOLD),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D9E2F3")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7FBFF")]),
    ]
    for i, r in enumerate(rows, start=1):
        if r["Status"] == "DO POTWIERDZENIA":
            table_style.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#FFF2CC")))
        if r["Status"] == "OPCJONALNIE":
            table_style.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#E2F0D9")))
    table.setStyle(TableStyle(table_style))
    story.append(table)

    doc.build(story)
    return path


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    OUTPUTS.mkdir(exist_ok=True)
    cat = catalog_map()
    hist = history_opp03()
    metrics = project_metrics()
    rows = apply_domain_rules(build_rows(cat), cat)
    issues = updated_issues()
    xlsx = write_xlsx(rows, issues, metrics, hist)
    pdf = write_pdf(rows, issues)
    print(f"XLSX={xlsx}")
    print(f"PDF={pdf}")


if __name__ == "__main__":
    main()
