import math
import sqlite3
import sys
import unicodedata
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


BASE = Path(r"C:\Users\Pawel Z\Documents\Codex\2026-05-28\files-mentioned-by-the-user-historia")
INPUTS = BASE / "inputs"
OUTPUTS = BASE / "outputs"
GPKG = INPUTS / "PW_Debe_Wielkie_OPP25.gpkg"
CATALOG = INPUTS / "KATALOGI PT_03_2026.xlsx"
HISTORY = INPUTS / "historia zamówień.xlsx"

ZLACZKA = "Złączka projektowana"
SPAW = "Spaw termiczny projektowany"


def register_fonts():
    fonts = Path(r"C:\Windows\Fonts")
    regular = fonts / "arial.ttf"
    bold = fonts / "arialbd.ttf"
    if regular.exists():
        pdfmetrics.registerFont(TTFont("Arial", str(regular)))
    if bold.exists():
        pdfmetrics.registerFont(TTFont("Arial-Bold", str(bold)))
    return ("Arial" if regular.exists() else "Helvetica", "Arial-Bold" if bold.exists() else "Helvetica-Bold")


FONT, FONT_BOLD = register_fonts()


SPECIAL = {
    ord("ł"): "l",
    ord("Ł"): "L",
    ord("ą"): "a",
    ord("Ą"): "A",
    ord("ć"): "c",
    ord("Ć"): "C",
    ord("ę"): "e",
    ord("Ę"): "E",
    ord("ń"): "n",
    ord("Ń"): "N",
    ord("ó"): "o",
    ord("Ó"): "O",
    ord("ś"): "s",
    ord("Ś"): "S",
    ord("ź"): "z",
    ord("Ź"): "Z",
    ord("ż"): "z",
    ord("Ż"): "Z",
}


def normalize(text):
    text = str(text).translate(SPECIAL)
    return "".join(c for c in unicodedata.normalize("NFKD", text) if not unicodedata.combining(c)).lower()


def fmt_sap(value):
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return ""
    return str(int(value)) if isinstance(value, float) else str(value)


def ceil_to(value, step):
    return int(math.ceil(float(value) / step) * step)


def catalog_map():
    cat = pd.read_excel(CATALOG, sheet_name="WYNIK")
    cat["Materiał SAP"] = pd.to_numeric(cat["Materiał SAP"], errors="coerce")

    preferences = {
        2200023907: "CORNING",
        2200023905: "CORNING",
        2200006863: "FCA",
        2200023710: "BELOS",
        2200028041: "BELOS",
        2200003277: "ROCK",
        2200022781: "YOFC",
        2200005414: "YOFC",
        2200005618: "FCA",
        2200003990: "FCA",
        2200003817: "LANDL",
        2200023869: "FCA",
        2200003716: "LANDL",
        2200003758: "LANDL",
        2200003756: "LANDL",
        2200028832: "CORNING",
    }

    out = {}
    text_columns = ["Nazwa", "Opis poz Dostawcy dł", "Opis długi prd.", "Nr poz Dostawcy w PZ", "Opis materiału SAP"]
    for sap, rows in cat.dropna(subset=["Materiał SAP"]).groupby("Materiał SAP"):
        sap_int = int(sap)
        selected = rows
        pref = preferences.get(sap_int)
        if pref:
            mask = (
                rows[text_columns]
                .fillna("")
                .astype(str)
                .agg(" ".join, axis=1)
                .str.upper()
                .str.contains(pref)
            )
            if mask.any():
                selected = rows.loc[mask]
        record = selected.iloc[0]
        out[sap_int] = {
            "sap": sap_int,
            "opis_sap": record.get("Opis materiału SAP", ""),
            "nr_dostawcy": record.get("Nr poz Dostawcy w PZ", ""),
            "opis_dostawcy": record.get("Opis poz Dostawcy dł", ""),
            "opis_dlugi": record.get("Opis długi prd.", ""),
            "dostawca": record.get("Nazwa", ""),
            "kategoria": record.get("podkategoria 2", ""),
            "jm_katalog": record.get("Jn zakupowa PZ", ""),
            "termin": record.get("Termin dostawy - dni", ""),
        }
    return out


def table_names(cur):
    names = [r[0] for r in cur.execute("select table_name from gpkg_contents")]

    def find(label):
        wanted = normalize(label)
        for name in names:
            if normalize(name) == wanted:
                return name
        for name in names:
            if wanted in normalize(name):
                return name
        raise KeyError(label)

    return {
        "lokale": find("Lokale"),
        "kable": find("Kable Swiatlowodowe"),
        "urzadzenia": find("Urzadzenia Pasywne"),
        "wlokna": find("Wlokna"),
        "zapasy": find("Zapasy"),
        "zestawienie": find("Zestawienie Czynnosci i Materialow"),
        "k_slup": find("K Slup"),
        "k_linia": find("K Linia Napowietrzna"),
    }


def project_metrics():
    con = sqlite3.connect(GPKG)
    con.row_factory = sqlite3.Row
    cur = con.cursor()
    t = table_names(cur)

    metrics = {}
    metrics["hh_by_node"] = [dict(r) for r in cur.execute(f"""
        select opp_osd, count(*) hh
        from [{t['lokale']}]
        group by opp_osd
        order by opp_osd
    """)]
    metrics["hh_total"] = sum(r["hh"] for r in metrics["hh_by_node"])

    metrics["devices"] = [dict(r) for r in cur.execute(f"""
        select wezel, typ_elementu, producent, model_urzadzenia, typ_obiektu, count(*) ilosc
        from [{t['urzadzenia']}]
        group by wezel, typ_elementu, producent, model_urzadzenia, typ_obiektu
        order by wezel, typ_elementu, model_urzadzenia
    """)]

    metrics["cables"] = [dict(r) for r in cur.execute(f"""
        with per_cable as (
          select odcinek_kabla, model_kabla, liczba_wlokien,
                 max(coalesce(dl_trasowa,0)) dl_trasowa,
                 max(coalesce(dl_optyczna,0)) dl_optyczna,
                 max(coalesce(dl_instalacyjna,0)) dl_instalacyjna
          from [{t['kable']}]
          group by odcinek_kabla, model_kabla, liczba_wlokien
        )
        select model_kabla, liczba_wlokien, count(*) liczba_odcinkow,
               round(sum(dl_trasowa), 2) trasa_m,
               round(sum(dl_optyczna), 2) optyczna_m,
               round(sum(dl_instalacyjna), 2) instalacyjna_m
        from per_cable
        group by model_kabla, liczba_wlokien
        order by liczba_wlokien, model_kabla
    """)]

    metrics["cable_edges"] = [dict(r) for r in cur.execute(f"""
        select odcinek_kabla, od, do, model_kabla, liczba_wlokien,
               max(coalesce(dl_instalacyjna,0)) dl_instalacyjna,
               max(coalesce(dl_trasowa,0)) dl_trasowa
        from [{t['kable']}]
        group by odcinek_kabla, od, do, model_kabla, liczba_wlokien
        order by odcinek_kabla
    """)]

    metrics["zapasy"] = [dict(r) for r in cur.execute(f"""
        select wezel, stan_stelazu, producent, model, count(*) ilosc, sum(dlugosc) dlugosc
        from [{t['zapasy']}]
        group by wezel, stan_stelazu, producent, model
        order by wezel, stan_stelazu
    """)]

    metrics["empty_material_schedule_count"] = cur.execute(f"select count(*) from [{t['zestawienie']}]").fetchone()[0]
    metrics["k_slup_count"] = cur.execute(f"select count(*) from [{t['k_slup']}]").fetchone()[0]
    metrics["k_linia_summary"] = [dict(r) for r in cur.execute(f"""
        select wlasciciel, count(*) odcinki, round(sum(coalesce(dl_instalacyjna,0)),2) dlugosc_m
        from [{t['k_linia']}]
        where dl_instalacyjna is not null
        group by wlasciciel
        order by wlasciciel
    """)]

    hh_by_node = {r["opp_osd"]: r["hh"] for r in metrics["hh_by_node"]}
    oap8_nodes = [
        r["wezel"]
        for r in metrics["devices"]
        if r["typ_elementu"] == "Przełącznica światłowodowa" and r["model_urzadzenia"] == "OAP_8"
    ]
    oap48_nodes = [
        r["wezel"]
        for r in metrics["devices"]
        if r["typ_elementu"] == "Przełącznica światłowodowa" and r["model_urzadzenia"] == "OAP_48"
    ]
    metrics["oap8_nodes"] = oap8_nodes
    metrics["oap48_nodes"] = oap48_nodes
    metrics["oap8_count"] = len(oap8_nodes)
    metrics["oap48_count"] = len(oap48_nodes)
    metrics["splitter_1x8_count"] = sum(
        r["ilosc"]
        for r in metrics["devices"]
        if r["typ_elementu"] == "Spliter" and "108" in str(r["model_urzadzenia"])
    )

    z_query = f"""
        with ends as (
          select coalesce(wezel_pocz,'') wezel, typ_polaczenia_pocz typ,
                 coalesce(pigtail_pocz_spaw,'') pigtail, typ_karty_pocz karta, odcinek_kabla cable
          from [{t['wlokna']}]
          union all
          select coalesce(wezel_kon,''), typ_polaczenia_kon,
                 coalesce(pigtail_kon_spaw,''), typ_karty_kon, odcinek_kabla
          from [{t['wlokna']}]
        )
        select wezel, karta, pigtail, count(*) n
        from ends
        where typ=? and cable is not null
        group by wezel, karta, pigtail
        order by wezel, karta, pigtail
    """
    metrics["z_endpoints"] = [dict(r) for r in cur.execute(z_query, (ZLACZKA,))]

    spaw_query = f"""
        with ends as (
          select coalesce(wezel_pocz,'') wezel, typ_polaczenia_pocz typ
          from [{t['wlokna']}]
          union all
          select coalesce(wezel_kon,''), typ_polaczenia_kon
          from [{t['wlokna']}]
        )
        select wezel, count(*) n
        from ends
        where typ=?
        group by wezel
        order by wezel
    """
    metrics["spaw_endpoints"] = [dict(r) for r in cur.execute(spaw_query, (SPAW,))]

    opp_nodes = oap48_nodes
    opp_pigtails = 0
    for node in opp_nodes:
        opp_pigtails += sum(r["n"] for r in metrics["z_endpoints"] if r["wezel"] == node)
    oap8_pigtails = sum(hh_by_node.get(node, 0) for node in oap8_nodes) + len(oap8_nodes)

    metrics["opp_pigtails"] = opp_pigtails
    metrics["oap8_pigtails"] = oap8_pigtails
    metrics["pigtails_total"] = opp_pigtails + oap8_pigtails
    metrics["adapters_opp"] = opp_pigtails
    metrics["spaw_endpoints_total"] = sum(r["n"] for r in metrics["spaw_endpoints"])
    metrics["splice_sleeves_order"] = ceil_to((metrics["pigtails_total"] + metrics["spaw_endpoints_total"]) * 1.10, 10)

    projected_stelaz_nodes = {
        r["wezel"]
        for r in metrics["zapasy"]
        if r["stan_stelazu"] == "Projektowany"
    }
    all_backbone_zapas_nodes = {
        r["wezel"]
        for r in metrics["zapasy"]
        if r["wezel"]
    }
    metrics["projected_stelaz_count"] = len(projected_stelaz_nodes)
    metrics["backbone_zapas_node_count"] = len(all_backbone_zapas_nodes)
    metrics["hook_count"] = metrics["backbone_zapas_node_count"] * 2
    metrics["uoz_5_7_count"] = 2
    metrics["uoo_10_count"] = metrics["hook_count"] - metrics["uoz_5_7_count"]

    con.close()
    return metrics


def history_summary():
    hist = pd.read_excel(HISTORY, sheet_name="Sheet1")
    hist["Ilosc_Zamówiona"] = pd.to_numeric(hist["Ilosc_Zamówiona"], errors="coerce").fillna(0)
    group = (
        hist.groupby(["NUMER \nINDEKSU\n(SAP)", "Nazwa Towaru", "JM"], dropna=False)
        .agg(ilosc=("Ilosc_Zamówiona", "sum"), wystapienia=("Nazwa Towaru", "size"))
        .reset_index()
        .sort_values("ilosc", ascending=False)
    )
    return group


def row(lp, status, category, sap, name, jm, qty_pw, qty_order, basis, notes="", catalog=None):
    c = catalog.get(sap, {}) if catalog and isinstance(sap, int) else {}
    return {
        "Lp": lp,
        "Status": status,
        "Kategoria": category,
        "SAP": fmt_sap(sap),
        "Nazwa materiału": name or c.get("opis_sap", ""),
        "JM": jm,
        "Ilość wg PW": qty_pw,
        "Ilość do zamówienia": qty_order,
        "Dostawca / pozycja katalogowa": f"{c.get('dostawca','')} | {c.get('nr_dostawcy','')}".strip(" |"),
        "Podstawa doboru": basis,
        "Uwagi": notes,
    }


def build_rows(catalog, m):
    rows = []

    def add(*args, **kwargs):
        rows.append(row(len(rows) + 1, *args, catalog=catalog, **kwargs))

    cable_12 = next(c for c in m["cables"] if c["liczba_wlokien"] == 12)
    cable_24 = next(c for c in m["cables"] if c["liczba_wlokien"] == 24)

    add(
        "DO POTWIERDZENIA",
        "Kable",
        2200022781,
        "KAB.ABON.ADSS.DROP FLEXTUBE 1X12 657A1",
        "m",
        f"{cable_12['instalacyjna_m']:.0f}",
        f"{ceil_to(cable_12['instalacyjna_m'], 10)}",
        "PW: ADSS LTC 12J, 1 odcinek OKH0002430/001, model 1x12, długość instalacyjna 260 m.",
        "Katalog nie ma pozycji opisanej dokładnie jako ADSS LTC 12J; dobrałem najbliższy 1x12 flex-tube. Potwierdzić klasę kabla/średnicę.",
    )
    add(
        "ZAMÓWIĆ",
        "Kable",
        2200003277,
        "KAB.ŚWIATŁ.ADSS 24J (6X4J) 657A1 3KN",
        "m",
        f"{cable_24['instalacyjna_m']:.0f}",
        f"{ceil_to(cable_24['instalacyjna_m'], 50)}",
        "PW: 9 odcinków ADSS LTC 24J 6x4; długość instalacyjna 1285 m.",
        "Pozycja katalogowa 24J 6x4; średnica ok. 10,8 mm, więc osprzęt 8-12/10-12 mm.",
    )

    add(
        "ZAMÓWIĆ",
        "Punkty pasywne",
        2200023907,
        "SKRZYNKA OAP-48",
        "szt",
        str(m["oap48_count"]),
        str(m["oap48_count"]),
        "PW: OPP0003 jako Corning OAP_48 na słupie.",
        "Wariant Corning z katalogu ma 48 adapterów SC/APC, tacki dla max 4 splitterów i 60 spawów.",
    )
    add(
        "ZAMÓWIĆ",
        "Punkty pasywne",
        2200023905,
        "SKRZYNKA OAP-8",
        "szt",
        str(m["oap8_count"]),
        str(m["oap8_count"]),
        "PW: 9 punktów OSD jako Corning OAP_8 na słupach.",
        "OAP_8 jest wyposażony w adaptery; pigtails trzeba domówić osobno.",
    )
    add(
        "ZAMÓWIĆ",
        "Punkty pasywne",
        2200023710,
        "STELAŻ ZAPASU KAB.DLA.OAP-8/24/48",
        "szt",
        str(m["projected_stelaz_count"]),
        str(m["projected_stelaz_count"]),
        "PW/Zapasy: 10 stelaży projektowanych Corning VOL_ST_50/10 przy OAP/OPP.",
        "",
    )
    add(
        "ZAMÓWIĆ",
        "Punkty pasywne",
        2200028041,
        "DYSTANS 150MM STELAŻA ZAPASU KAB OAP",
        "szt",
        f"{m['projected_stelaz_count']} stelaży x 2",
        str(m["projected_stelaz_count"] * 2),
        "Przyjęto 2 dystanse 150 mm na stelaż zapasu OAP na słupie.",
        "Jeśli wybrany komplet stelaża dostawcy zawiera dystanse, tę pozycję skorygować po potwierdzeniu dostawy.",
    )

    add(
        "ZAMÓWIĆ",
        "Optyka",
        2200006863,
        "SPLITTER PLC ALUBOX 1X8/657/900/SCA",
        "szt",
        str(m["splitter_1x8_count"]),
        str(m["splitter_1x8_count"]),
        "PW: 3 splittery S-PL-108-TUBE-900-SCA w OPP i po 1 w każdym z 9 OAP_8.",
        "Model zgodny z projektem: FCA S-PL-108-TUBE-900-SCA.",
    )
    add(
        "ZAMÓWIĆ",
        "Optyka",
        2200005414,
        "PIGTAIL 657 SC/APC DŁ.3M",
        "szt",
        f"{m['opp_pigtails']} OPP + {m['oap8_pigtails']} OAP",
        str(m["pigtails_total"]),
        "OPP: wszystkie zakończenia kablowe z projektu = 17. OAP_8: HH na OSD + dosył = 60 + 9.",
        "W OAP adaptery są w wyposażeniu, ale pigtaile domawiamy.",
    )
    add(
        "DO POTWIERDZENIA",
        "Optyka",
        2200005618,
        "ADAPTER SC/APC",
        "szt",
        str(m["adapters_opp"]),
        str(m["adapters_opp"]),
        "Zasada OPP: adaptery w liczbie zakończeń OPP = 17.",
        "Uwaga: katalogowa OAP-48 Corning ma 48 adapterów w wyposażeniu. Pozycję zamawiać tylko, jeśli inwestor wymaga oddzielnego adaptera dla OPP mimo wyposażenia OAP-48.",
    )
    add(
        "ZAMÓWIĆ",
        "Optyka",
        2200003990,
        "OSŁONKA SPAWU ŚWIATŁOWODU OS-45",
        "szt",
        f"{m['pigtails_total']} pigtaili + {m['spaw_endpoints_total']} spawów włókno-włókno",
        str(m["splice_sleeves_order"]),
        "Pigtail = 1 spaw; dodatkowo warstwa Włókna pokazuje 26 projektowanych spawów włókno-włókno. Dodano ok. 10% zapasu i zaokrąglono do 10 szt.",
        "",
    )

    add(
        "ZAMÓWIĆ",
        "Napowietrzne",
        2200003716,
        "HAK UNIWERSALNY",
        "szt",
        f"{m['backbone_zapas_node_count']} słupów/węzłów x 2",
        str(m["hook_count"]),
        "Zasada wykonawcza: na jeden słup zamawiamy dwa haki; policzono 11 słupów/węzłów z warstwy Zapasy dla kabla dosyłowego/backbone.",
        "Nie użyłem warstwy K Słup, bo zawiera szerszy zakres koncepcyjny.",
    )
    add(
        "ZAMÓWIĆ",
        "Napowietrzne",
        2200003817,
        "UCHWYT ODCIĄGOWY UOO/10-W 10-12MM",
        "szt",
        "ADSS 24J 6x4",
        str(m["uoo_10_count"]),
        "Uchwyty odciągowe dla ADSS 24J 6x4 o średnicy ok. 10,8 mm; liczba uchwytów razem z 5-7 mm = liczba haków.",
        "20 szt. obejmuje 18 końców odcinków 24J i 2 szt. zapasu na montaż.",
    )
    add(
        "DO POTWIERDZENIA",
        "Napowietrzne",
        2200023869,
        "UCHWYT ODCIĄGOWY KLINOWY DLA 4.8-6.2MM",
        "szt",
        "ADSS LTC 12J 1x12",
        str(m["uoz_5_7_count"]),
        "Dla dobranego kabla 12J flex-tube ok. 5-6 mm: 2 końce odcinka OKH0002430/001.",
        "Nie dotyczy abonenckich ADSS 2J, których nie ujmuję w zamówieniu.",
    )
    add(
        "ZAMÓWIĆ",
        "Napowietrzne",
        2200003758,
        "TAŚMA STALOWA TSM/20-07-J",
        "rol",
        "montaż OAP/stelaży/haków",
        "1",
        "Osprzęt do mocowania elementów na słupach.",
        "Przyjąłem jedną rolkę 50 m, jeżeli brak zapasu magazynowego.",
    )
    add(
        "ZAMÓWIĆ",
        "Napowietrzne",
        2200003756,
        "KLAMRA DO TAŚMY STALOWEJ TSK/20-J",
        "pak",
        "komplet z taśmą",
        "1",
        "Klamry do taśmy stalowej 20 mm.",
        "Przyjąłem jedno opakowanie, jeżeli brak zapasu magazynowego.",
    )
    return rows


def build_issues(m):
    drop = next(c for c in m["cables"] if c["liczba_wlokien"] == 2)
    return [
        {
            "Temat": "Abonencki ADSS 2J",
            "Opis": f"PW ma {drop['liczba_odcinkow']} odcinków ADSS LTC 2J, trasa {drop['trasa_m']:.1f} m, ale tylko {drop['instalacyjna_m']:.1f} m instalacyjnie. Na zadaniach FTTH SI nie zamawiamy zapasów abonenckich kabli napowietrznych.",
            "Rekomendacja": "Nie ujmuję ADSS 2J ani uchwytów 5-7 mm dla przyłączy abonenckich. Do sprawdzenia zostaje pojedynczy odcinek z instalacyjną 38 m.",
        },
        {
            "Temat": "ADSS LTC 12J",
            "Opis": "Projekt opisuje kabel jako ADSS LTC 12J 1x12. W katalogu nie znalazłem pozycji nazwanej dokładnie ADSS LTC 12J; najbliższa pozycja to ADSS DROP FLEXTUBE 1x12.",
            "Rekomendacja": "Potwierdzić z projektantem/inwestorem, czy dla odcinka OKH0002430/001 ma być kabel 1x12 flex-tube 5-6 mm, czy mocniejszy ADSS 12J 3kN 3x4.",
        },
        {
            "Temat": "Adaptery w OPP OAP-48",
            "Opis": "Reguła wykonawcza wymaga adapterów w OPP, ale katalogowa skrzynka Corning OAP-48 jest już wyposażona w 48 adapterów SC/APC.",
            "Rekomendacja": "W tabeli zostawiłem 17 adapterów jako DO POTWIERDZENIA, żeby nie przeoczyć reguły OPP, ale prawdopodobnie nie trzeba ich domawiać przy wariancie Corning OAP-48.",
        },
        {
            "Temat": "Warstwy K_*",
            "Opis": f"Warstwa K Słup ma {m['k_slup_count']} słupów, a K Linia Napowietrzna obejmuje także Dębe Wielkie/Chroślę/Bykowiznę i szerszy zakres koncepcyjny.",
            "Rekomendacja": "Nie liczyłem haków/przelotówek z całej warstwy K_*. Haki/odciągi policzyłem z węzłów z warstwy Zapasy dla projektowanych kabli 12J/24J.",
        },
        {
            "Temat": "Uchwyty przelotowe ADSS",
            "Opis": "Dla kabli 12J/24J mogą być potrzebne uchwyty przelotowe na podporach pośrednich. Warstwa kabli pokazuje odcinki między OPP/OSD, natomiast K Linia Napowietrzna nie daje jednoznacznego filtra na samo OPP25.",
            "Rekomendacja": "Nie wpisałem przelotówek do zamówienia liczbowo. Do ustalenia z trasą montażową: UP-JP 8-12 mm dla ADSS 24J i ewentualnie 5-8 mm dla 12J flex-tube.",
        },
        {
            "Temat": "Zestawienie Czynności i Materiałów",
            "Opis": f"Gotowa tabela zestawienia w GPKG ma {m['empty_material_schedule_count']} rekordów.",
            "Rekomendacja": "BOM został policzony z warstw kabli, urządzeń, włókien i zapasów oraz zweryfikowany katalogiem PT_03_2026.",
        },
    ]


def write_xlsx(rows, issues, metrics, history):
    path = OUTPUTS / "lista_materialow_Debe_Wielkie_OPP25.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Zamówienie"
    headers = list(rows[0].keys())
    ws.append(headers)
    for r in rows:
        ws.append([r[h] for h in headers])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_cells in ws.iter_rows(min_row=2):
        status = row_cells[1].value
        fill = PatternFill("solid", fgColor="FFF2CC") if status == "DO POTWIERDZENIA" else None
        for cell in row_cells:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(top=thin, left=thin, right=thin, bottom=thin)
            if fill:
                cell.fill = fill
    widths = [5, 17, 18, 12, 38, 8, 18, 18, 42, 56, 62]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wi = wb.create_sheet("Do wyjaśnienia")
    wi.append(["Temat", "Opis", "Rekomendacja"])
    for issue in issues:
        wi.append([issue["Temat"], issue["Opis"], issue["Rekomendacja"]])
    for cell in wi[1]:
        cell.fill = PatternFill("solid", fgColor="C65911")
        cell.font = header_font
    for col, width in enumerate([34, 90, 90], start=1):
        wi.column_dimensions[get_column_letter(col)].width = width
    for row_cells in wi.iter_rows():
        for cell in row_cells:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    wm = wb.create_sheet("Przedmiar PW")
    wm.append(["Sekcja", "Dane"])
    for item in metrics["hh_by_node"]:
        wm.append(["HH", str(item)])
    for item in metrics["cables"]:
        wm.append(["Kable", str(item)])
    for item in metrics["devices"]:
        wm.append(["Urządzenia", str(item)])
    for item in metrics["zapasy"]:
        wm.append(["Zapasy", str(item)])
    for item in metrics["z_endpoints"]:
        wm.append(["Złącza włókien", str(item)])
    for item in metrics["spaw_endpoints"]:
        wm.append(["Spawy włókien", str(item)])
    for item in metrics["k_linia_summary"]:
        wm.append(["K Linia - informacyjnie", str(item)])
    wm.column_dimensions["A"].width = 26
    wm.column_dimensions["B"].width = 150
    for row_cells in wm.iter_rows():
        for cell in row_cells:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    wh = wb.create_sheet("Historia ogółem")
    wh.append(list(history.columns))
    for _, record in history.head(200).iterrows():
        wh.append([record[c] for c in history.columns])
    for col in range(1, len(history.columns) + 1):
        wh.column_dimensions[get_column_letter(col)].width = 28
    for cell in wh[1]:
        cell.fill = PatternFill("solid", fgColor="70AD47")
        cell.font = header_font

    wb.save(path)
    return path


def para(text, style):
    return Paragraph(str(text).replace("\n", "<br/>"), style)


def write_pdf(rows, issues):
    path = OUTPUTS / "lista_materialow_Debe_Wielkie_OPP25.pdf"
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="BodyPL", fontName=FONT, fontSize=7.2, leading=8.4, alignment=TA_LEFT))
    styles.add(ParagraphStyle(name="SmallPL", fontName=FONT, fontSize=6.2, leading=7.2, alignment=TA_LEFT))
    styles.add(ParagraphStyle(name="TitlePL", fontName=FONT_BOLD, fontSize=15, leading=18, spaceAfter=6))
    styles.add(ParagraphStyle(name="HeadingPL", fontName=FONT_BOLD, fontSize=10, leading=12, spaceBefore=8, spaceAfter=4))

    doc = SimpleDocTemplate(
        str(path),
        pagesize=landscape(A4),
        rightMargin=8 * mm,
        leftMargin=8 * mm,
        topMargin=8 * mm,
        bottomMargin=8 * mm,
    )
    story = []
    story.append(Paragraph("Lista materiałów do zamówienia - PW Dębe Wielkie OPP25", styles["TitlePL"]))
    story.append(Paragraph(f"Data opracowania: {datetime.now().strftime('%Y-%m-%d %H:%M')}. Źródła: GPKG PW, katalog PT_03_2026, historia zamówień.", styles["BodyPL"]))
    story.append(Spacer(1, 4))

    story.append(Paragraph("Punkty do wyjaśnienia", styles["HeadingPL"]))
    issue_data = [[para("Temat", styles["SmallPL"]), para("Opis", styles["SmallPL"]), para("Rekomendacja", styles["SmallPL"])]]
    for issue in issues:
        issue_data.append([para(issue["Temat"], styles["SmallPL"]), para(issue["Opis"], styles["SmallPL"]), para(issue["Rekomendacja"], styles["SmallPL"])])
    issue_table = Table(issue_data, colWidths=[42 * mm, 116 * mm, 120 * mm], repeatRows=1)
    issue_table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, 0), FONT_BOLD),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#C65911")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D9E2F3")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#FFF7EA")]),
    ]))
    story.append(issue_table)
    story.append(PageBreak())

    story.append(Paragraph("Tabela zamówienia", styles["HeadingPL"]))
    headers = ["Lp", "Status", "Kategoria", "SAP", "Nazwa materiału", "JM", "Ilość do zamówienia", "Podstawa doboru", "Uwagi"]
    data = [[para(h, styles["SmallPL"]) for h in headers]]
    for r in rows:
        data.append([
            para(r["Lp"], styles["SmallPL"]),
            para(r["Status"], styles["SmallPL"]),
            para(r["Kategoria"], styles["SmallPL"]),
            para(r["SAP"], styles["SmallPL"]),
            para(r["Nazwa materiału"], styles["SmallPL"]),
            para(r["JM"], styles["SmallPL"]),
            para(r["Ilość do zamówienia"], styles["SmallPL"]),
            para(r["Podstawa doboru"], styles["SmallPL"]),
            para(r["Uwagi"], styles["SmallPL"]),
        ])
    table = Table(data, colWidths=[8 * mm, 25 * mm, 24 * mm, 20 * mm, 48 * mm, 10 * mm, 22 * mm, 72 * mm, 49 * mm], repeatRows=1)
    table_style = [
        ("FONTNAME", (0, 0), (-1, 0), FONT_BOLD),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D9E2F3")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7FBFF")]),
    ]
    for i, r in enumerate(rows, start=1):
        if r["Status"] == "DO POTWIERDZENIA":
            table_style.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#FFF2CC")))
    table.setStyle(TableStyle(table_style))
    story.append(table)
    doc.build(story)
    return path


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    OUTPUTS.mkdir(exist_ok=True)
    cat = catalog_map()
    metrics = project_metrics()
    rows = build_rows(cat, metrics)
    issues = build_issues(metrics)
    hist = history_summary()
    xlsx = write_xlsx(rows, issues, metrics, hist)
    pdf = write_pdf(rows, issues)
    print(f"XLSX={xlsx}")
    print(f"PDF={pdf}")


if __name__ == "__main__":
    main()
